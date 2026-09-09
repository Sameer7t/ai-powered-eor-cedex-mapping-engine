import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(BASE_DIR)

def generate_excel_report(invoices: list[dict], output_path: str):
    print(f"\n--- [EXCEL WRITER] Initializing workbook for {len(invoices)} invoices ---")

    if not invoices:
        print(f"--- [EXCEL WRITER] WARNING: No invoices provided for {output_path}. Aborting write. ---")
        return

    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estimates"

    header_font = Font(name="Aptos Display", bold=True)
    body_font = Font(name="Aptos Display")
    header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

    headers = [
        "Container Number", "Container Type", "Repair Company", "Pool",
        "Job Description", "LOCN", "CMP", "DMG", "RPR",
        "LENGTH", "WIDTH", "QTY", "MAN.HRS", "HRS.COST", "MATRL.Cost",
        "TOTAL", "Grand Total",
    ]

    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_index, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    current_row = 2
    total_jobs_written = 0

    for invoice in invoices:
        container_number = invoice.get("container_id", "UNKNOWN")
        first_row_of_block = current_row

        jobs = invoice.get("jobs", [])
        if not jobs:
            print(f"--- [EXCEL WRITER] Skipping {container_number}: 'jobs' list is empty. ---")
            continue

        print(f"--- [EXCEL WRITER] Writing Container {container_number} to Rows {current_row} - {current_row + len(jobs) - 1} ---")

        for job in jobs:
            is_first_row = (current_row == first_row_of_block)

            # Safely resolve costs, coercing to float to prevent openpyxl crashes
            labour_cost_val = float(job.get("labour_cost") or job.get("labour") or job.get("service_cost") or 0.0)
            mat_cost_val = float(job.get("material_cost_aed") or job.get("material_cost") or 0.0)

            row_values = {
                "A": container_number if is_first_row else None,
                "B": invoice.get("container_type", "") if is_first_row else None,
                "C": invoice.get("depot_name", "") if is_first_row else None,
                "D": "",
                "E": str(job.get("job_description", "")),
                "F": str(job.get("location", "")),
                "G": str(job.get("component", "")),
                "H": str(job.get("damage", "")),
                "I": str(job.get("repair", "")),
                "J": 1, "K": 1, "L": 1,
                "M": float(job.get("manhour") or 0.0),
                "N": labour_cost_val,
                "O": mat_cost_val,
            }

            for col_letter, value in row_values.items():
                cell = ws[f"{col_letter}{current_row}"]
                cell.value = value
                cell.font = body_font

            total_cell = ws[f"P{current_row}"]
            total_cell.value = f"=N{current_row}+O{current_row}"
            total_cell.font = body_font

            for col_idx in range(1, 18):
                ws.cell(row=current_row, column=col_idx).border = thin_border

            current_row += 1
            total_jobs_written += 1

        last_row = current_row - 1
        if last_row >= first_row_of_block:
            grand_total_cell = ws[f"Q{last_row}"]
            grand_total_cell.value = f"=SUM(P{first_row_of_block}:P{last_row})"
            grand_total_cell.font = body_font

        current_row += 1

    for col_index, header in enumerate(headers, start=1):
        column_letter = get_column_letter(col_index)
        ws.column_dimensions[column_letter].width = len(str(header)) + 4

    print(f"--- [EXCEL WRITER] Attempting to save {total_jobs_written} rows to disk... ---")

    try:
        wb.save(output_path)
        print(f"--- [EXCEL WRITER] SUCCESS: File committed to {output_path} ---")
    except Exception as e:
        print(f"--- [EXCEL WRITER] CRITICAL SYSTEM ERROR DURING SAVE: {e} ---")